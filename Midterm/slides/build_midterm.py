"""Build the 10/28 mid-term presentation (python-pptx): demo-centred, 12 slides.

    python slides/build_midterm.py      # -> slides/RouteVision_期中報告_315833033.pptx

Figures in slides/assets/midterm/ are screenshots of the web app (python app/server.py, ?demo=0&step=...).
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
A = os.path.join(HERE, 'assets', 'midterm')
OUT = os.path.join(HERE, 'RouteVision_期中報告_315833033.pptx')
FONT = 'Microsoft JhengHei'

rgb = lambda h: RGBColor.from_string(h)
INK, BODY, MUTED, FOOT = rgb('1C1E21'), rgb('33373C'), rgb('5A5F66'), rgb('6B6F75')
TEAL, TEALD, TEAL_SOFT = rgb('0F766E'), rgb('0B5D56'), rgb('E4F1EE')
PAPER, PAPER2, LINE = rgb('F6F5F1'), rgb('FCFBF8'), rgb('E3DFD6')
DARK, DARK2, DARK3 = rgb('13201E'), rgb('1F302D'), rgb('2A4440')
MINT, LIGHT, LIGHT2, LIGHT3 = rgb('6FD3C3'), rgb('F4F8F6'), rgb('D7E4E0'), rgb('A9BFBA')
RED, BAR, BAR_EDA, STRIKE = rgb('D9474C'), rgb('9CBAB4'), rgb('C9C3B6'), rgb('9A958B')

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, M = 13.333, 0.6
CW = W - 2 * M


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15, after=0):
    """paras: str, or list of paragraphs; a paragraph is str or list of (text, {size, color, bold})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(o.get('size', size))
            r.font.bold = o.get('bold', bold)
            r.font.color.rgb = o.get('color', color)
    return tb


def box(slide, x, y, w, h, fill, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def pic(slide, name, x, y, w, h=None):
    p = slide.shapes.add_picture(os.path.join(A, name), Inches(x), Inches(y), Inches(w), Inches(h) if h else None)
    return p


def title(slide, t, y=0.55, size=32):
    text(slide, M, y, CW, 0.8, t, size=size, bold=True)


def footer(slide, n):
    text(slide, M, 6.95, 6, 0.3, 'RouteVision · 林孟霆', size=11, color=FOOT)
    text(slide, W - M - 1, 6.95, 1, 0.3, str(n), size=11, color=FOOT, align=PP_ALIGN.RIGHT)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def new(color):
    s = prs.slides.add_slide(BLANK)
    bg(s, color)
    return s


def table(slide, x, y, w, rows, widths, size=15, header_color=MUTED, highlight=None, align_right_from=1, row_h=0.42):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for j, cw in enumerate(widths):
        t.columns[j].width = Inches(w * cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = TEAL_SOFT if i == highlight else (PAPER if i == 0 else PAPER2)
            c.margin_left = c.margin_right = Inches(0.08)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if j >= align_right_from else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = str(v)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = i == 0 or i == highlight
            r.font.color.rgb = header_color if i == 0 else (TEALD if i == highlight else BODY)
    return t


# ============================================================ 1 cover
s = new(DARK)
text(s, 0.8, 1.95, 7, 0.4, 'MMIP 期中專題 · 鑑別式 AI 系統', size=15, color=MINT, bold=True)
text(s, 0.8, 2.4, 7.2, 1.1, 'RouteVision', size=60, color=LIGHT, bold=True)
text(s, 0.8, 3.65, 6.8, 1.3, ['placement 一完成，', '就預測 routing 後哪裡會塞'], size=26, color=LIGHT2, spacing=1.25)
text(s, 0.8, 6.3, 8, 0.4, '林孟霆 · 315833033 · 交大電子所 · 2026/10/28', size=14, color=LIGHT3)
HOT = {(1, 1), (1, 2), (1, 3), (1, 4), (1, 5), (2, 0), (2, 4), (2, 5), (3, 4), (4, 4)}
WARM = {(2, 1), (2, 3), (3, 0), (3, 5), (4, 3), (5, 4)}
for r in range(8):
    for c in range(8):
        col = RED if (r, c) in HOT else DARK3 if (r, c) in WARM else DARK2
        box(s, 8.45 + c * 0.48, 1.6 + r * 0.48, 0.42, 0.42, col, radius=0.12)
notes(s, '（約 15 秒）大家好，我是林孟霆。我的期中專題叫 RouteVision：在晶片實體設計中，placement 一做完，就預測 routing 之後哪裡會繞線壅塞。右邊這個 8×8 的格子就是輸出：紅色是預測會塞的區塊。')

# ============================================================ 2 recap
s = new(PAPER2)
title(s, '回顧：placement 一完成，就預測哪裡會塞')
cards = [('問題', ['routing 要跑數小時才知道哪裡塞', 'EDA 內建粗估（RUDY、eGR）和真實結果的相關係數只有 0.55、0.65']),
         ('做法', ['輸入 4 張 placement 特徵圖＋6 個設計參數', '輸出 8×8 熱點、整張使用率圖；比較 RF、MLP、CNN、Multimodal、ViT、Ensemble、U-Net']),
         ('評估', ['CircuitNet 2.0（14 nm）：GPU、AI 加速器、RISC-V，各 96 張版圖', '三個設計輪流當「沒看過的晶片」，閾值只在驗證集決定'])]
cw = (CW - 0.5) / 3
for k, (h, ps) in enumerate(cards):
    x = M + k * (cw + 0.25)
    box(s, x, 1.6, cw, 3.25, PAPER, LINE)
    text(s, x + 0.3, 1.85, cw - 0.6, 0.45, h, size=18, color=TEAL, bold=True)
    text(s, x + 0.3, 2.45, cw - 0.6, 2.3, ps, size=15, color=BODY, spacing=1.3, after=10)
box(s, M, 5.2, CW, 0.95, DARK, radius=0.12)
text(s, M + 0.35, 5.2, 1.5, 0.95, '期中新增', size=17, color=MINT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
x = 2.55
for label, pw in (('U-Net 逐像素預測', 2.35), ('可信度', 1.25), ('Grad-CAM 判斷依據', 2.45), ('四步驟網頁', 1.65)):
    box(s, x, 5.43, pw, 0.5, DARK2, radius=0.5)
    text(s, x, 5.43, pw, 0.5, label, size=14, color=LIGHT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    x += pw + 0.2
footer(s, 2)
notes(s, '（約 30 秒）快速回顧構想報告：routing 要跑好幾個小時才知道哪裡塞，EDA 工具內建的粗估又不夠準。我用 placement 後就有的特徵圖和設計參數，預測哪裡會塞，並比較多種方法。評估一律以「沒看過的晶片」為準，三個設計輪流當測試。期中新增的是下面這四項，今天的重點就是 demo 這些新功能。')

# ============================================================ 3 results
s = new(PAPER)
title(s, '結果：沒看過的晶片上，AI 更準也更穩定')
rows = [('U-Net 多模態', 0.984, 0.013, 'best'), ('Ensemble', 0.971, 0.014, ''), ('U-Net', 0.970, 0.020, ''),
        ('Random Forest', 0.967, 0.020, ''), ('Multimodal', 0.966, 0.018, ''), ('ViT', 0.955, 0.021, ''),
        ('CNN', 0.951, 0.022, ''), ('MLP', 0.931, 0.047, ''), ('eGR（EDA）', 0.911, 0.069, 'eda'), ('RUDY（EDA）', 0.881, 0.061, 'eda')]
for i, (name, v, sd, kind) in enumerate(rows):
    y = 1.55 + i * 0.44
    c = TEALD if kind == 'best' else FOOT if kind == 'eda' else BODY
    text(s, M, y, 2.0, 0.34, name, size=14, color=c, bold=kind == 'best', align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    bw = (v - 0.85) / 0.15 * 4.9
    box(s, 2.8, y + 0.05, bw, 0.24, TEAL if kind == 'best' else BAR_EDA if kind == 'eda' else BAR, radius=0.15)
    text(s, 2.8 + bw + 0.12, y, 1.6, 0.34, f'{v:.3f} ± {sd:.3f}', size=12, color=c, bold=kind == 'best', anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.95, 6.0, 7.5, 0.3, 'ROC-AUC，三個設計輪流當測試的平均 ± 標準差；橫軸從 0.85 開始', size=11, color=FOOT)
stats = [('3 / 3', '個設計都是第一，包括 eGR 最強的 openc910-1（0.994 vs 0.987）'),
         ('191 / 288', '張版圖的熱點 F1 贏過 eGR'),
         ('± 0.013', '標準差只有 eGR（± 0.069）的五分之一，換晶片也穩定'),
         ('0.937', 'PR-AUC：標出的熱點有多可信（eGR 0.718）')]
for i, (big, lab) in enumerate(stats):
    y = 1.45 + i * 1.18
    text(s, 9.0, y, 3.8, 0.55, big, size=28, color=TEAL, bold=True)
    text(s, 9.0, y + 0.55, 3.75, 0.6, lab, size=13, color=BODY, spacing=1.2)
footer(s, 3)
notes(s, '（約 55 秒）這是最主要的結果：三個設計輪流當沒看過的晶片，ROC-AUC 的平均和標準差。所有學習方法平均都贏過兩個 EDA 粗估。最好的是 U-Net 多模態，0.984，三個設計都是第一；連 eGR 最強的 openc910-1，它也從 0.987 被超過。逐張版圖比較，288 張裡有 191 張的熱點 F1 贏過 eGR。而且標準差只有 eGR 的五分之一，代表換一顆晶片表現也很穩定。最後是 PR-AUC，它看的是模型說「這裡會塞」的時候可不可信，亂猜的分數等於熱點比例，熱點少時比 ROC-AUC 嚴格：U-Net 多模態是 0.937，eGR 只有 0.718，差距主要來自 eGR 的大量誤報。')

# ============================================================ 4 system
s = new(PAPER)
title(s, '系統：把分析變成四個步驟的故事')
p = pic(s, 'ui_hot.png', M, 1.55, 7.7)
p.line.color.rgb = rgb('D8D3C8')
items = [('1', '一步只講一件事', '熱點 → 使用率圖 → 可信度 → 判斷依據'),
         ('2', '同一個版面', 'EDA ↔ AI ↔ 真實三圖並排，滑鼠同步'),
         ('3', '先給結論', '一句話摘要，細節收進展開'),
         ('4', '不讓人等', '預先計算：第一次分析 62 秒 → 0.2 秒')]
for i, (n, h, d) in enumerate(items):
    y = 1.65 + i * 1.08
    box(s, 8.75, y, 0.42, 0.42, TEAL, shape=MSO_SHAPE.OVAL)
    text(s, 8.75, y, 0.42, 0.42, n, size=15, color=LIGHT, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.35, y - 0.02, 3.4, 0.45, h, size=18, bold=True)
    text(s, 9.35, y + 0.42, 3.4, 0.6, d, size=13, color=MUTED, spacing=1.2)
footer(s, 4)
notes(s, '（約 30 秒）系統是 Flask 後端加 HTML/JavaScript 前端。UI 的設計理念是把分析變成四個步驟的故事：每一步只講一件事，而且版面都一樣，左邊是 EDA 粗估、中間是 AI、右邊是真實結果，上方先給一句話結論。因為 ViT 在 CPU 上推論要 20 秒，server 啟動時會先把範例都算好，第一次分析從 62 秒降到 0.2 秒。接下來直接 demo。')

# ============================================================ 5 demo divider
s = new(DARK)
text(s, 0.8, 1.35, 6, 0.4, 'LIVE DEMO', size=15, color=MINT, bold=True)
text(s, 0.8, 1.8, 11.5, 0.9, '一張沒看過的晶片，五個操作', size=40, color=LIGHT, bold=True)
steps = [('①', '熱點預測', 'AI 比粗估少誤報'), ('②', '使用率圖', '亮點一：U-Net'), ('③', '可信度', '亮點二：不確定度'),
         ('④', '判斷依據', 'Grad-CAM'), ('⑤', '上傳新版圖', '沒有答案的設計')]
sw = (W - 1.6 - 4 * 0.2) / 5
for i, (n, h, d) in enumerate(steps):
    x = 0.8 + i * (sw + 0.2)
    box(s, x, 3.15, sw, 1.65, DARK2, radius=0.1)
    text(s, x + 0.25, 3.32, sw - 0.4, 0.35, n, size=16, color=MINT, bold=True)
    text(s, x + 0.25, 3.72, sw - 0.4, 0.45, h, size=20, color=LIGHT, bold=True)
    text(s, x + 0.25, 4.22, sw - 0.4, 0.4, d, size=13, color=LIGHT3)
text(s, 0.8, 5.2, 11.5, 0.4, '按 → 換步驟、↑↓ 換版圖。接下來 5 頁是每一步要看的重點（現場設備有問題時照著講）', size=15, color=LIGHT3)
text(s, 0.8, 6.3, 11.5, 0.4, '備用影片：[Demo 影片連結]', size=15, color=LIGHT3)
notes(s, '（轉場，約 5 秒）接下來直接操作網頁。\n報告前準備：先執行 python app/server.py，等左側範例旁的綠點全亮（約 3 分鐘，之後重開會讀快取），瀏覽器開 http://127.0.0.1:5000 。\n現場：切到瀏覽器，依 ①→⑤ 操作；每一步要講的重點在後面 5 頁的講稿。若設備有問題，就直接翻後面 5 頁，或播放備用影片。')


# ============================================================ 6-9 demo steps
def step_slide(n, color, card_color, eyebrow, op, head, img, callouts, note):
    s = new(color)
    text(s, M, 0.45, 6, 0.35, eyebrow, size=14, color=TEAL, bold=True)
    text(s, 6.6, 0.47, W - M - 6.6, 0.35, op, size=12, color=FOOT, align=PP_ALIGN.RIGHT)
    text(s, M, 0.82, CW, 0.7, head, size=28, bold=True)
    iw = 10.6
    pic(s, img, (W - iw) / 2, 1.65, iw)
    cw = (CW - 0.5) / 3
    for k, runs in enumerate(callouts):
        x = M + k * (cw + 0.25)
        box(s, x, 5.65, cw, 0.95, card_color, radius=0.12)
        text(s, x + 0.25, 5.65, cw - 0.5, 0.95, [runs], size=15, color=BODY, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    footer(s, n)
    notes(s, note)


B = lambda t: (t, {'bold': True, 'color': TEALD})
step_slide(6, PAPER2, PAPER, 'DEMO ① 熱點預測', '操作：點左側第一張 Vortex-small', 'AI 比 EDA 粗估少了很多誤報', 'crop_hot.png',
           [[('左：eGR 抓到 7/10，', {}), B('誤報 8 格')],
            [('中：U-Net 抓到 9/10，', {}), B('只誤報 1 格')],
            [('上方一句話結論：', {}), B('F1 0.900 vs 0.560')]],
           '（約 25 秒）\n操作：點左側「沒看過的晶片」第一張 Vortex-small，停在第 ① 步。\n要講：這張晶片完全沒參與訓練。三張圖並排，左邊是 EDA 工具的 eGR，中間是 AI，右邊是 routing 後的真實熱點。紅色是抓到、斜線是誤報、藍框是漏掉。eGR 抓到 7 格，但誤報了 8 格；U-Net 多模態抓到 9 格，只誤報 1 格。上方這句話就是結論，F1 從 0.56 提升到 0.90。\n可補充：右上角下拉選單可以換成其他 AI 方法比較；滑鼠移到任何一格，三張圖會同步標示並顯示數值。')
step_slide(7, PAPER, PAPER2, 'DEMO ② 使用率圖 · 亮點一', '操作：按 →', 'U-Net 直接畫出 routing 後的整張使用率圖', 'crop_map.png',
           [[('這張：相關係數 ', {}), B('0.738'), ('，eGR 0.604', {})],
            [('三設計平均：', {}), B('0.81'), (' vs 0.65；SSIM 0.45 vs 0.24', {})],
            [('畫出左下的低使用率方塊和水平繞線條紋', {})]],
           '（約 30 秒）\n操作：按右方向鍵到第 ② 步。\n要講：第一個亮點。前面只回答「這格是不是熱點」，U-Net 直接預測 routing 後整張使用率圖，越亮越擠。和右邊的真實結果比，U-Net 畫出了左下角低使用率的方塊、還有水平的繞線條紋；eGR 則比較雜亂。這張的相關係數是 0.738，eGR 是 0.604；三個設計平均是 0.81 對 0.65。\n技術補充（被問再講）：編碼器是 ImageNet 預訓練的 ResNet-18，設計參數用 FiLM 注入；消融實驗顯示加入設計參數後，區塊熱點 ROC-AUC 從 0.970 升到 0.984，三個設計都提升。')
step_slide(8, PAPER2, PAPER, 'DEMO ③ 可信度 · 亮點二', '操作：按 →', '模型標出自己沒把握的地方', 'crop_trust.png',
           [[('這張：錯 6 格，', {}), B('4 格在紫色框裡'), ('（共 7 格）', {})],
            [('平均：最沒把握的 20%，涵蓋 ', {}), B('78%'), (' 的錯誤', {})],
            [('openc910-1：', {}), B('0.93'), ('，簡單對照組只有 0.61', {})]],
           '（約 30 秒）\n操作：按右方向鍵到第 ③ 步。\n要講：第二個亮點。Ensemble 有三個做法完全不同的模型：Random Forest、Multimodal、ViT。同一格如果三個模型給的機率差很多，就代表沒把握，中間紫色框就是這些區塊。這不需要答案，推論時就能算。這張版圖 Ensemble 錯了 6 格，其中 4 格就在它標出的 7 個紫色框裡。三個設計平均，只要人工檢查最沒把握的 20% 區塊，就能抓到 78% 的錯誤。\n技術補充（被問再講）：不確定度＝三個模型機率的標準差。對照組是「分數離閾值越近越不確定」，在分布差最多的 openc910-1 上，對照組的 AUROC 只有 0.61，模型分歧有 0.93。')
step_slide(9, PAPER, PAPER2, 'DEMO ④ 判斷依據', '操作：按 →；右上可換 CNN、Multimodal', 'Grad-CAM：模型看哪裡判斷熱點', 'crop_why.png',
           [[('暖色＝判斷熱點時', {}), B('最依賴'), ('的區域', {})],
            [('最熱的地方落在', {}), B('真實熱點'), ('（虛線框）上', {})],
            [('左下也有一塊關注，但不是熱點', {})]],
           '（約 20 秒）\n操作：按右方向鍵到第 ④ 步，預設是 U-Net 多模態。\n要講：Grad-CAM 用梯度算出每個特徵對答案的重要度，再畫回原圖。暖色是模型判斷熱點時最依賴的區域。最熱的地方落在真實熱點的虛線框上，代表模型學到的是真正的壅塞樣式；左下也有一塊關注、但不是熱點，這也是模型可以改進的地方。\n可補充：右上可以切到 CNN、Multimodal 比較。Random Forest、MLP、ViT 沒有卷積特徵圖，不能用 Grad-CAM。U-Net 用的是分割版的 Seg-Grad-CAM。')

# ============================================================ 10 upload
s = new(PAPER2)
text(s, M, 0.45, 6, 0.35, 'DEMO ⑤ 上傳新版圖', size=14, color=TEAL, bold=True)
text(s, 6.6, 0.47, W - M - 6.6, 0.35, '檔案：demo/upload_examples/', size=12, color=FOOT, align=PP_ALIGN.RIGHT)
text(s, M, 0.82, CW, 0.7, '還沒 routing 的新設計，也能直接預測', size=28, bold=True)
up = [[('按左上「上傳」', {})],
      [('選 ', {}), ('vortex_D_no_answer.npz', {'bold': True}), ('：只有特徵圖和設計參數，沒有答案', {})],
      [('四個步驟照樣能看：熱點、使用率圖、沒把握的區塊、判斷依據', {})],
      [('右邊的「真實結果」顯示無資料，因為還沒 routing', {})]]
for i, runs in enumerate(up):
    y = 1.9 + i * 1.0
    box(s, M, y, 0.42, 0.42, TEAL, shape=MSO_SHAPE.OVAL)
    text(s, M, y, 0.42, 0.42, str(i + 1), size=15, color=LIGHT, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M + 0.6, y - 0.02, 5.2, 0.9, [runs], size=16, color=BODY, spacing=1.25)
table(s, 7.0, 1.9, 5.73, [('檔案', '熱點', 'U-Net F1', 'eGR F1'), ('vortex_A', 8, '0.778', '0.583'), ('vortex_B', 15, '0.889', '0.765'),
                          ('vortex_C', 21, '0.606', '0.529'), ('vortex_D', 17, '無答案', '–')],
      [0.34, 0.18, 0.24, 0.24], size=15, highlight=4, row_h=0.48)
text(s, 7.0, 4.45, 5.73, 0.9, '全部取自沒參與訓練的 Vortex-small，依熱點數挑選，沒看模型表現。已預先算好，上傳約 0.6 秒出結果。',
     size=13, color=MUTED, spacing=1.3)
footer(s, 10)
notes(s, '（約 20 秒）\n操作：按左上「上傳」，選 demo/upload_examples/vortex_D_no_answer.npz。\n要講：最後模擬真實情境：一個剛做完 placement、還沒 routing 的新設計，檔案裡沒有答案。系統一樣能預測熱點、畫出使用率圖、標出沒把握的區塊，右邊的真實結果則顯示沒有資料。\n備用：如果老師想看有答案的上傳，可以傳 A、B、C。C 是 U-Net 只抓到一半熱點的例子，可以誠實說明單張版圖會有變動；整體表現要看三個設計的平均。')

# ============================================================ 11 capacity
s = new(PAPER2)
title(s, '模型夠大嗎？瓶頸在資料的多樣性')
table(s, M, 1.6, 7.6, [('模型', '參數', '同設計', '沒看過的晶片', '落差'), ('CNN', '0.3 M', '0.994', '0.951', '−0.043'),
                       ('Multimodal', '0.3 M', '0.995', '0.966', '−0.029'), ('ViT-B/16', '86 M', '0.993', '0.955', '−0.038'),
                       ('U-Net 多模態', '14 M', '0.997', '0.984', '−0.013')],
      [0.3, 0.15, 0.17, 0.22, 0.16], size=16, highlight=4, row_h=0.55)
text(s, M, 4.45, 7.6, 0.3, 'ROC-AUC，三個設計平均。同設計＝和訓練同設計的驗證版圖', size=11, color=FOOT)
box(s, 8.6, 1.6, 4.13, 2.75, PAPER, LINE)
text(s, 8.9, 1.8, 3.6, 2.4, [[('同設計大家都 ≈ 0.99：', {'bold': True}), ('0.3 M 的小 CNN 就夠，題目太簡單', {})],
                             [('大 300 倍的 ViT：', {'bold': True}), ('換晶片沒有比較好', {})],
                             [('逐像素＋設計參數：', {'bold': True, 'color': TEALD}), ('落差最小，換晶片最穩', {'color': TEALD})]],
     size=15, color=BODY, spacing=1.3, after=12)
text(s, M, 5.3, CW, 0.5, [[('限制　', {'bold': True, 'color': BODY}), ('只有 3 個設計、每組只訓練一次（seed 42）；「多擠才算熱點」的閾值仍需依設計調整。', {})]],
     size=15, color=MUTED)
footer(s, 11)
notes(s, '（約 40 秒）大家可能會問：模型是不是太小？在同設計的驗證集上，0.3 M 參數的 CNN 就已經 0.99，題目本身太簡單，看不出差異。大 300 倍的 ViT 換一顆晶片也沒有比較好。真正讓跨晶片表現變好的，是把任務改成逐像素預測、再加上設計參數，落差從 0.03 到 0.04 降到 0.013。所以瓶頸是資料的多樣性，而不是模型容量。限制是只有 3 個設計、每組只訓練一次，而且熱點門檻還需要依設計調整，這是未來工作。')

# ============================================================ 12 summary
s = new(PAPER)
title(s, '總結')
points = [[('用「沒看過的晶片」評估：U-Net 多模態 ROC-AUC ', {}), ('0.984 ± 0.013', {'bold': True}), ('，三個設計都贏 EDA 粗估', {})],
          [('不只判斷熱點：預測整張使用率圖、標出沒把握的區塊、說明判斷依據', {})],
          [('從資料前處理到網頁的完整鑑別式 AI 系統，用上課程 Week 1–6 的工具', {})]]
for i, runs in enumerate(points):
    y = 1.65 + i * 1.05
    text(s, M, y, 0.5, 0.5, str(i + 1), size=22, color=TEAL, bold=True)
    text(s, M + 0.6, y, CW - 0.6, 0.9, [runs], size=20, color=INK, spacing=1.3)
box(s, M, 5.0, CW, 0.95, PAPER2, LINE)
text(s, M + 0.35, 5.0, 1.5, 0.95, '未來工作', size=16, color=TEAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, M + 1.85, 5.0, CW - 2.2, 0.95, '加入更多 CircuitNet 設計　·　依設計自動調整熱點門檻　·　期末：LLM Agent 解讀結果，給工程師修改建議',
     size=14, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 12)
notes(s, '（約 20 秒）總結三點：第一，用沒看過的晶片評估，U-Net 多模態在三個設計上都勝過 EDA 工具的粗估；第二，系統不只判斷熱點，還能預測整張使用率圖、標出模型沒把握的地方，並說明判斷依據；第三，這是一個從資料到網頁的完整系統。未來會加入更多設計、讓熱點門檻能依設計自動調整，期末再加上 LLM Agent 解讀結果、給工程師修改建議。謝謝大家。')

prs.save(OUT)
print('wrote', OUT.encode('ascii', 'replace').decode())
