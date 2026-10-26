"""Build the 10/14 mid-term idea presentation (python-pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = r'C:\Users\yp455\Downloads\MMIP_project\slides'
A = HERE + r'\assets\\'
FONT = 'Microsoft JhengHei'

NAVY = RGBColor(0x16, 0x1D, 0x33)
NAVY2 = RGBColor(0x24, 0x2E, 0x4D)
AMBER = RGBColor(0xF2, 0xA3, 0x0F)
ORANGE = RGBColor(0xE4, 0x57, 0x2E)
INK = RGBColor(0x1F, 0x26, 0x3A)
MUTED = RGBColor(0x5B, 0x64, 0x78)
CARD = RGBColor(0xF2, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD6, 0xDB, 0xE4)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: str, or list of paragraphs; each paragraph is str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.space_after = Pt(spacing)
        segs = p if isinstance(p, list) else [(p, {})]
        for seg, o in segs:
            r = para.add_run()
            r.text = seg
            r.font.name = FONT
            r.font.size = Pt(o.get('size', size))
            r.font.bold = o.get('bold', bold)
            r.font.color.rgb = o.get('color', color)
    return tb


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def badge(slide, x, y, label, d=0.5, fill=AMBER, color=NAVY):
    c = box(slide, x, y, d, d, fill, MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = color


def title(slide, num, t, sub=None):
    badge(slide, 0.6, 0.55, num)
    text(slide, 1.3, 0.45, 11.4, 0.7, t, size=30, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        text(slide, 1.3, 1.12, 11.4, 0.4, sub, size=15, color=MUTED)


def footer(slide, page, dark=False):
    text(slide, 0.6, 7.0, 9, 0.3, '林孟霆 315833033｜MMIP 期中 Milestone 構想報告', size=10,
         color=RGBColor(0x9A, 0xA3, 0xB5) if dark else MUTED)
    text(slide, 11.7, 7.0, 1.0, 0.3, str(page), size=10, align=PP_ALIGN.RIGHT,
         color=RGBColor(0x9A, 0xA3, 0xB5) if dark else MUTED)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
s.shapes.add_picture(A + 'label_util.png', Inches(8.35), Inches(1.2), Inches(4.4), Inches(4.4))
text(s, 8.35, 5.75, 4.4, 0.6, '實際資料：一個晶片版圖的繞線使用率分布（CircuitNet-N14，亮處為繞線擁擠）',
     size=11, color=RGBColor(0xB8, 0xC0, 0xD0))
text(s, 0.8, 1.35, 7.2, 0.5, 'MMIP 期中 Milestone 構想報告｜鑑別式 AI 系統', size=16, color=AMBER, bold=True)
text(s, 0.8, 1.95, 7.3, 1.9, [[('RouteVision', {'size': 48, 'bold': True, 'color': WHITE})],
                              [('晶片版圖繞線壅塞預測系統', {'size': 34, 'bold': True, 'color': WHITE})]])
text(s, 0.8, 4.05, 7.0, 0.9, '在 placement 階段，用影像辨識提早找出 routing 壅塞熱點，並解釋模型看到了什麼',
     size=17, color=RGBColor(0xCF, 0xD6, 0xE4))
text(s, 0.8, 5.55, 7.0, 0.9, [[('林孟霆  315833033', {'size': 16, 'bold': True, 'color': WHITE})],
                              [('交大電子所 平行運算系統實驗室（AI-aided HLS）', {'size': 14, 'color': RGBColor(0xB8, 0xC0, 0xD0)})]])
notes(s, '大家好，我是林孟霆，來自電子所平行運算系統實驗室，研究方向是 AI 輔助的高階合成。'
         '我的題目是 RouteVision，一個用影像辨識預測晶片版圖繞線壅塞的鑑別式 AI 系統。右邊這張就是真實資料中一個晶片的壅塞分布，亮的地方就是繞線擁擠的區域。（約 20 秒）')

# ---------------------------------------------------------------- 2. Problem & system
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, '1', '系統想解決的問題：壅塞發現得太晚', '預計開發的系統與要解決的問題')
box(s, 0.6, 1.8, 5.9, 4.8, CARD)
text(s, 0.95, 2.05, 5.3, 0.5, '現況的痛點', size=20, bold=True, color=ORANGE)
text(s, 0.95, 2.65, 5.3, 3.8, [
    [('晶片設計流程：', {'bold': True}), ('placement → routing → 檢查', {})],
    [('真正的壅塞', {'bold': True}), ('要等 routing 跑完才確定；placement 階段只有粗估（RUDY、early global routing），實測和最終結果的相關係數只有約 0.40／0.57', {})],
    [('一旦壅塞：', {'bold': True}), ('要回頭修改設計再重跑，迭代成本高；HLS 產生的硬體在 FPGA 上也有同樣問題', {})],
], size=16, spacing=14)
box(s, 0.95, 5.25, 5.2, 1.05, WHITE)
text(s, 1.2, 5.35, 4.7, 0.85, [[('目標：', {'bold': True, 'color': ORANGE}),
                               ('在 placement 後就做出比現有粗估更準的壅塞預測', {'bold': True, 'color': NAVY})]],
     size=15, anchor=MSO_ANCHOR.MIDDLE)
box(s, 6.85, 1.8, 5.9, 4.8, NAVY)
text(s, 7.2, 2.05, 5.3, 0.5, 'RouteVision 的做法', size=20, bold=True, color=AMBER)
steps = [('輸入', 'placement 後的版圖特徵圖（影像）'),
         ('判斷', '每個區塊會不會壅塞（熱點分類）'),
         ('輸出', '壅塞熱度圖 + Grad-CAM 解釋'),
         ('效益', '比 RUDY／eGR 粗估更準，提早修正')]
for i, (k, v) in enumerate(steps):
    y = 2.75 + i * 0.92
    badge(s, 7.2, y, str(i + 1), d=0.45)
    text(s, 7.85, y - 0.02, 4.7, 0.5, [[(k + '：', {'bold': True, 'color': AMBER}), (v, {'color': WHITE})]],
         size=16, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)
notes(s, '晶片設計要先擺放元件，再繞線。真正的壅塞要等 routing 跑完才確定；placement 階段 EDA 工具只有 RUDY 和 early global routing 這類粗估，我用實際資料算過，它們和最終結果的相關係數只有約 0.4 到 0.57。一旦壅塞就得回頭改設計重跑，非常耗時。'
         '這和我的研究也直接相關：HLS 產生的硬體在 FPGA 上也常因為壅塞造成 timing 不過。'
         '所以我想做一個系統，在 placement 之後用影像辨識做出比這些粗估更準的壅塞預測，並用 Grad-CAM 解釋模型的依據。（約 45 秒）')

# ---------------------------------------------------------------- 3. Data
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, '2', '資料：CircuitNet 開源 EDA 資料集', '預計使用的資料類型與資料來源')
imgs = [('in_density.png', '輸入：cell 密度'), ('in_rudy.png', '輸入：繞線需求 RUDY'),
        ('in_pin.png', '輸入：腳位繞線需求'), ('label_util.png', '標註：實際繞線使用率')]
for i, (f, cap) in enumerate(imgs):
    x = 0.6 + i * 1.95
    s.shapes.add_picture(A + f, Inches(x), Inches(1.85), Inches(1.75), Inches(1.75))
    text(s, x, 3.7, 1.75, 0.35, cap, size=12, color=ORANGE if i == 3 else MUTED, align=PP_ALIGN.CENTER,
         bold=(i == 3))
text(s, 0.6, 4.25, 7.6, 0.4, '已實際下載 Vortex-small 檢查：96 個樣本，每張約 400–460 格見方',
     size=14, color=INK, bold=True)
text(s, 0.6, 4.75, 7.6, 1.9, [
    '資料類型：晶片版圖切成格子的多通道「影像」，每格代表 placement 後的特徵',
    '來源：北京大學等單位公開的 CircuitNet（BSD-3 授權，Hugging Face 下載）',
    '只需 routability 特徵；先用 3 個小設計（下載約 3.5 GB，解壓後約 10 倍，只解壓需要的特徵並縮成 256×256）',
], size=14, spacing=8, color=MUTED)
stats = [('10,345', '個版圖樣本'), ('8', '個開源設計\nRISC-V／GPU／ML 加速器'), ('14 nm', 'FinFET 製程')]
for i, (big, small) in enumerate(stats):
    y = 1.85 + i * 1.6
    box(s, 8.75, y, 4.0, 1.4, CARD)
    text(s, 9.0, y + 0.12, 3.6, 0.7, big, size=32, bold=True, color=NAVY)
    text(s, 9.0, y + 0.78, 3.6, 0.6, small, size=12, color=MUTED)
footer(s, 3)
notes(s, '資料使用 CircuitNet，是北京大學等單位專為機器學習輔助 EDA 公開的資料集，BSD-3 授權。'
         'N14 版本有一萬多個版圖樣本，來自 RISC-V、GPU、ML 加速器等 8 個設計。'
         '我已經下載其中最小的 Vortex-small 檢查過：每個樣本是約四百多格見方的陣列，尺寸隨 floorplan 設定不同，輸入是 cell 密度和繞線需求，標註是實際繞線後的壅塞程度，就是最右邊這張。'
         '只需要其中 25 GB 的繞線特徵，先用 3 個小設計，Colab 就能處理。（約 50 秒）')

# ---------------------------------------------------------------- 4. Course techniques
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, '3', '預計使用的課程技術', '期中比較不同的鑑別式方法，期末加入 Agent')
cards = [('傳統方法與 baseline', 'EDA 粗估（RUDY、eGR）；OpenCV\n特徵 ＋ Random Forest／SVM', 'Week 1–2'),
         ('MLP＋多模態', 'MLP baseline；影像 CNN ＋\n設計參數 MLP 的雙輸入模型', 'Week 2–3'),
         ('CNN／U-Net', 'CNN 判斷熱點；U-Net 預測\n整張壅塞熱度圖', 'Week 3'),
         ('Vision Transformer', '微調 ViT，比較與 CNN\n在版圖影像上的差異', 'Week 4'),
         ('可解釋性與評估', 'Grad-CAM 標出依據的區域\nROC-AUC、Macro-F1、混淆矩陣', 'Week 2–3'),
         ('期末：Agentic AI', 'LLM Agent ＋ RAG（壅塞成因\n與修正方法）＋ Function Call', '期末')]
for i, (h, b, wk) in enumerate(cards):
    col, row = i % 3, i // 3
    x, y = 0.6 + col * 4.1, 1.85 + row * 2.45
    last = i == 5
    box(s, x, y, 3.85, 2.2, NAVY if last else CARD)
    text(s, x + 0.3, y + 0.25, 2.4, 0.45, h, size=18, bold=True, color=AMBER if last else NAVY)
    text(s, x + 2.65, y + 0.3, 0.95, 0.35, wk, size=11, bold=True, align=PP_ALIGN.RIGHT,
         color=AMBER if last else ORANGE)
    text(s, x + 0.3, y + 0.9, 3.3, 1.1, b, size=14, color=WHITE if last else MUTED)
footer(s, 4)
notes(s, '技術上會用到課程每週的內容：先用 EDA 工具內建的粗估、以及 OpenCV 特徵加上 Random Forest 當 baseline，再比較 MLP、CNN 和 ViT，U-Net 則用來預測整張壅塞圖。另外做一個多模態模型：每個樣本都有時脈頻率、floorplan 使用率等設計參數，把影像和這些參數一起輸入，對應老師示範中的 multi-head 架構。'
         '評估用 ROC-AUC 和 Macro-F1，並用 Grad-CAM 解釋模型看的是哪些區域。期末再加入 LLM Agent、RAG 和 Function Call。（約 40 秒）')

# ---------------------------------------------------------------- 5. Architecture
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, '4', '系統流程與技術架構', '期中：可操作的網頁 demo；期末：在同一套系統上加入 Agent')
flow = [('資料', 'CircuitNet\n版圖特徵圖\n＋ 設計參數'), ('前處理', 'resize 256×256\n正規化、切區塊\n依設計切分'),
        ('模型比較', 'EDA 粗估／RF\nMLP／CNN／ViT\n多模態／U-Net'), ('輸出', '熱點分類\n壅塞熱度圖\nGrad-CAM'),
        ('網頁 Demo', 'Streamlit\n上傳版圖\n比較各方法')]
bw, gap = 2.0, 0.53
for i, (h, b) in enumerate(flow):
    x = 0.6 + i * (bw + gap)
    box(s, x, 1.9, bw, 2.35, NAVY if i == 4 else CARD)
    text(s, x + 0.2, 2.05, bw - 0.4, 0.45, h, size=17, bold=True, color=AMBER if i == 4 else NAVY,
         align=PP_ALIGN.CENTER)
    text(s, x + 0.15, 2.6, bw - 0.3, 1.5, b, size=13, color=WHITE if i == 4 else MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        arr = box(s, x + bw + 0.08, 2.9, gap - 0.16, 0.35, AMBER, MSO_SHAPE.RIGHT_ARROW)
# final-stage agent band
box(s, 0.6, 4.65, 12.15, 2.05, CARD, line=LINE)
text(s, 0.95, 4.85, 3.0, 0.45, '期末延伸：Agentic AI', size=17, bold=True, color=ORANGE)
text(s, 0.95, 5.35, 2.9, 1.2, '使用者上傳版圖並提問\n「哪裡會壅塞？怎麼改？」', size=13, color=MUTED)
tools = [('Function Call', '呼叫壅塞預測\n與 Grad-CAM'), ('RAG', '壅塞成因與\n修正方法知識庫'),
         ('Guardrail', '信心不足時\n建議人工確認'), ('未來延伸', '移植到 FPGA（需自建\nVivado 資料）接 HLS')]
for i, (h, b) in enumerate(tools):
    x = 4.05 + i * 2.17
    box(s, x, 4.9, 1.97, 1.6, WHITE, line=LINE)
    text(s, x + 0.12, 5.0, 1.73, 0.4, h, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x + 0.1, 5.45, 1.77, 0.95, b, size=12, color=MUTED, align=PP_ALIGN.CENTER)
footer(s, 5)
notes(s, '系統流程是：讀入版圖特徵圖和設計參數，前處理時統一大小、正規化、切成區塊，並且依設計切分訓練和測試集。'
         '接著比較五種方法，輸出熱點分類、壅塞熱度圖和 Grad-CAM，最後做成 Streamlit 網頁讓使用者上傳版圖、比較各方法。'
         '期末在同一套系統上加入 Agent，透過 Function Call 呼叫預測模型、用 RAG 查壅塞成因和修正方法；未來如果自行產生 FPGA 的資料，也希望把方法接回我的 HLS 研究。（約 50 秒）')

# ---------------------------------------------------------------- 6. Outcomes & risks
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, '5', '預期成果與可能遇到的問題', '已從實際資料發現的挑戰，以及對應方法')
text(s, 0.6, 1.8, 4.3, 0.45, '預期成果', size=20, bold=True, color=NAVY)
outs = [('可操作網頁', '上傳版圖，比較 5 種方法\n的預測結果與 Grad-CAM'),
        ('量化比較', '和 EDA 粗估及 5 種方法\n比較 ROC-AUC、Macro-F1'),
        ('泛化驗證', '在「沒看過的設計」上測試\n檢驗是否真的實用')]
for i, (h, b) in enumerate(outs):
    y = 2.45 + i * 1.4
    badge(s, 0.6, y, str(i + 1), d=0.45)
    text(s, 1.25, y - 0.05, 3.9, 1.2, [[(h, {'bold': True, 'color': INK, 'size': 16})], [(b, {'color': MUTED, 'size': 14})]])
risks = [('標籤太稀疏', '繞線 overflow 只出現在約 1% 的格子', '改用繞線使用率門檻定義熱點'),
         ('資料洩漏', '同一設計的樣本彼此非常相似', '依設計切分訓練／測試集'),
         ('不是一般照片', '多通道數值圖，ViT 預訓練權重不一定適用', '比較從頭訓練與微調'),
         ('尺寸與數量', '每張尺寸不同；Vortex-small 只有 96 筆', '統一縮放並切區塊；加入其他設計')]
text(s, 5.4, 1.8, 7.3, 0.45, '可能遇到的問題 → 對策', size=20, bold=True, color=NAVY)
for i, (h, why, fix) in enumerate(risks):
    y = 2.4 + i * 1.12
    box(s, 5.4, y, 7.35, 0.98, CARD)
    text(s, 5.65, y + 0.1, 2.1, 0.8, h, size=15, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.7, y + 0.1, 4.9, 0.8, [[(why, {'color': MUTED, 'size': 13})], [('→ ' + fix, {'color': INK, 'size': 14, 'bold': True})]],
         anchor=MSO_ANCHOR.MIDDLE)
footer(s, 6)
notes(s, '預期成果有三個：可操作的網頁、和 EDA 粗估及五種方法的量化比較，以及在沒看過的設計上測試泛化能力。'
         '我已經從實際資料發現幾個挑戰：繞線 overflow 只出現在約 1% 的格子，太稀疏，所以會改用繞線使用率的門檻定義熱點；'
         '同一個設計的樣本非常相似，必須依設計切分，否則分數會虛高；輸入是多通道數值圖不是照片，會比較從頭訓練和微調；每張尺寸不同而且 Vortex-small 只有 96 筆，會統一縮放、切成區塊，並加入其他設計。（約 50 秒）')

# ---------------------------------------------------------------- 7. Plan / closing
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, 0.8, 0.8, 11.5, 0.8, '時程規劃', size=32, bold=True, color=WHITE)
plan = [('10/14', '構想報告', '題目、資料、架構確定\n已完成資料下載與檢查'),
        ('10/28', '期中 Milestone', '前處理、5 種方法比較\nGrad-CAM、Streamlit 網頁'),
        ('期末', 'Final Project', 'LLM Agent、RAG\nFunction Call、Guardrail')]
for i, (d, h, b) in enumerate(plan):
    x = 0.8 + i * 4.05
    box(s, x, 2.0, 3.75, 3.0, NAVY2)
    text(s, x + 0.35, 2.25, 3.1, 0.7, d, size=30, bold=True, color=AMBER)
    text(s, x + 0.35, 3.0, 3.1, 0.5, h, size=18, bold=True, color=WHITE)
    text(s, x + 0.35, 3.6, 3.1, 1.2, b, size=14, color=RGBColor(0xCF, 0xD6, 0xE4))
    if i < 2:
        box(s, x + 3.8, 3.3, 0.2, 0.35, AMBER, MSO_SHAPE.RIGHT_ARROW)
text(s, 0.8, 5.6, 11.7, 0.8, '謝謝大家，歡迎提問與建議', size=24, bold=True, color=WHITE)
footer(s, 7, dark=True)
notes(s, '時程上，今天確定題目和資料，資料已經下載檢查完；10/28 完成前處理、五種方法比較、Grad-CAM 和網頁；期末再加入 Agent 和 RAG。謝謝大家。（約 20 秒）')

prs.save(HERE + r'\RouteVision_期中構想報告_v3.pptx')
print('saved')




